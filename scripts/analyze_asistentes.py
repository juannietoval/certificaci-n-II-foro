# -*- coding: utf-8 -*-
import openpyxl
import re

file_path = r'c:\Users\Lenovo\Desktop\Juan\Registro de Asistencia - II Foro de Editores de Revistas Científicas 2026 (respuestas).xlsx'
wb = openpyxl.load_workbook(file_path)
ws = wb.active

records = []
for r in range(2, ws.max_row + 1):
    timestamp = ws.cell(r, 1).value
    email = str(ws.cell(r, 2).value or '').strip()
    nombre = str(ws.cell(r, 3).value or '').strip()
    tipo_doc = str(ws.cell(r, 4).value or '').strip()
    doc_raw = str(ws.cell(r, 5).value or '').strip()
    institucion = str(ws.cell(r, 6).value or '').strip()
    pais_ciudad = str(ws.cell(r, 7).value or '').strip()
    rol = str(ws.cell(r, 8).value or '').strip()
    
    if doc_raw.endswith('.0'):
        doc_raw = doc_raw[:-2]
        
    records.append({
        'row': r,
        'timestamp': timestamp,
        'email': email,
        'nombre': nombre,
        'tipo_doc': tipo_doc,
        'doc': doc_raw,
        'institucion': institucion,
        'pais_ciudad': pais_ciudad,
        'rol': rol
    })

print(f'Total records in spreadsheet: {len(records)}')

# 1. Duplicate emails
email_counts = {}
for rec in records:
    em = rec['email'].lower()
    email_counts[em] = email_counts.get(em, 0) + 1

dup_emails = {k: v for k, v in email_counts.items() if v > 1 and k != ''}
print(f'\n--- DUPLICATE EMAILS ({len(dup_emails)}) ---')
for em, cnt in dup_emails.items():
    matching = [r for r in records if r['email'].lower() == em]
    print(f'  {em} ({cnt} times):')
    for m in matching:
        print(f'    Row {m["row"]}: Name="{m["nombre"]}", Doc="{m["doc"]}", Rol="{m["rol"]}"')

# 2. Duplicate docs
doc_counts = {}
for rec in records:
    d = rec['doc']
    if d and d != 'None':
        doc_counts[d] = doc_counts.get(d, 0) + 1

dup_docs = {k: v for k, v in doc_counts.items() if v > 1}
print(f'\n--- DUPLICATE DOC NUMBERS ({len(dup_docs)}) ---')
for d, cnt in dup_docs.items():
    matching = [r for r in records if r['doc'] == d]
    print(f'  Doc {d} ({cnt} times):')
    for m in matching:
        print(f'    Row {m["row"]}: Email={m["email"]}, Name="{m["nombre"]}", Rol="{m["rol"]}"')

# 3. Duplicate names
name_counts = {}
for rec in records:
    nm = re.sub(r'\s+', ' ', rec['nombre']).lower()
    name_counts[nm] = name_counts.get(nm, 0) + 1

dup_names = {k: v for k, v in name_counts.items() if v > 1}
print(f'\n--- DUPLICATE NAMES ({len(dup_names)}) ---')
for nm, cnt in dup_names.items():
    matching = [r for r in records if re.sub(r'\s+', ' ', r['nombre']).lower() == nm]
    print(f'  Name "{nm}" ({cnt} times):')
    for m in matching:
        print(f'    Row {m["row"]}: Email={m["email"]}, Doc="{m["doc"]}", Rol="{m["rol"]}"')

# 4. Check for anomalous names (e.g. emails in name field)
print(f'\n--- ANOMALOUS NAMES ---')
for rec in records:
    if '@' in rec['nombre'] or len(rec['nombre']) < 4 or rec['nombre'].isdigit():
        print(f'  Row {rec["row"]}: Name="{rec["nombre"]}", Email="{rec["email"]}", Doc="{rec["doc"]}"')
