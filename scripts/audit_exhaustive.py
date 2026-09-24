import os
import glob
import json
import openpyxl

def main():
    root = r'c:\Users\Lenovo\Desktop\Juan'
    
    # Locate exact source files
    prog_files = glob.glob(os.path.join(root, '*Programacion*.xlsx'))
    asist_files = glob.glob(os.path.join(root, '*Registro de Asistencia*.xlsx'))
    
    print(f"Programacion file: {prog_files}")
    print(f"Asistencia file: {asist_files}")
    
    wb_prog = openpyxl.load_workbook(prog_files[0])
    wb_asist = openpyxl.load_workbook(asist_files[0])
    
    print("\n--- PROGRAMACION SHEETS ---")
    for name in wb_prog.sheetnames:
        ws = wb_prog[name]
        print(f"Sheet '{name}': {ws.max_row} rows, {ws.max_column} cols")
        
    print("\n--- ASISTENCIA SHEETS ---")
    for name in wb_asist.sheetnames:
        ws = wb_asist[name]
        print(f"Sheet '{name}': {ws.max_row} rows, {ws.max_column} cols")
        # Print header
        headers = [cell.value for cell in ws[1]]
        print(f"Headers: {headers}")

    # Load master registro
    master_path = os.path.join(root, 'sistema_certificados', 'registro_certificados_foro_2026.xlsx')
    wb_master = openpyxl.load_workbook(master_path)
    print(f"\n--- MASTER REGISTRO SHEETS ---")
    for name in wb_master.sheetnames:
        ws = wb_master[name]
        print(f"Sheet '{name}': {ws.max_row} rows, {ws.max_column} cols")
        headers = [cell.value for cell in ws[1]]
        print(f"Headers: {headers}")

    # Load participantes.json
    json_path = os.path.join(root, 'sistema_certificados', 'data', 'participantes.json')
    with open(json_path, 'r', encoding='utf-8') as f:
        participantes = json.load(f)
    print(f"\nTotal in participantes.json: {len(participantes)}")
    ponentes = [p for p in participantes if p.get('rol', '').upper() == 'PONENTE']
    asistentes = [p for p in participantes if p.get('rol', '').upper() == 'ASISTENTE']
    print(f"Ponentes: {len(ponentes)}, Asistentes: {len(asistentes)}")

if __name__ == '__main__':
    main()
