# -*- coding: utf-8 -*-
import os
import pymupdf
import openpyxl

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"

dir_ponentes = os.path.join(BASE, "certificados", "ponentes")
dir_asistentes = os.path.join(BASE, "certificados", "asistentes")
dir_output_pdf = os.path.join(BASE, "output", "pdf")
dir_output_preview = os.path.join(BASE, "output", "preview")

files_pon = os.listdir(dir_ponentes)
files_asi = os.listdir(dir_asistentes)
files_out_pdf = os.listdir(dir_output_pdf)
files_out_prev = os.listdir(dir_output_preview)

print(f"Certificados Ponentes: {len(files_pon)} files")
print(f"Certificados Asistentes: {len(files_asi)} files")
print(f"Output PDFs (Web portal): {len(files_out_pdf)} files")
print(f"Output Previews (Web portal): {len(files_out_prev)} files")

# Check all PDFs in asistentes
corrupt = []
multipage = []
small_files = []

for f in files_asi:
    p = os.path.join(dir_asistentes, f)
    sz = os.path.getsize(p)
    if sz < 100000:
        small_files.append((f, sz))
    try:
        doc = pymupdf.open(p)
        if len(doc) != 1:
            multipage.append((f, len(doc)))
        doc.close()
    except Exception as e:
        corrupt.append((f, str(e)))

print("\n--- AUDIT RESULTS FOR ASISTENTES PDFS ---")
print(f"Total checked: {len(files_asi)}")
print(f"Corrupted: {len(corrupt)}")
print(f"Multi-page (>1 page): {len(multipage)}")
print(f"Suspiciously small (<100KB): {len(small_files)}")

# Check Excel
excel_path = os.path.join(BASE, "registro_certificados_foro_2026.xlsx")
wb = openpyxl.load_workbook(excel_path)
print(f"\n--- EXCEL AUDIT ({os.path.basename(excel_path)}) ---")
print(f"Sheets: {wb.sheetnames}")
for sname in wb.sheetnames:
    ws = wb[sname]
    print(f"  Sheet '{sname}': {ws.max_row - 1} data rows, {ws.max_column} columns")
