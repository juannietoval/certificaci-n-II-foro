import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"
json_path = os.path.join(BASE, "data", "participantes.json")

with open(json_path, "r", encoding="utf-8") as f:
    participantes = json.load(f)

asistentes = [p for p in participantes if p.get("rol") == "ASISTENTE"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Envio_Asistentes"
ws.views.sheetView[0].showGridLines = True

headers = [
    "Código Certificado",
    "Nombre Asistente",
    "Tratamiento",
    "Correo Electrónico",
    "Documento Oficial",
    "Institución",
    "País / Ciudad",
    "Asunto Correo",
    "Enlace Validación Web (QR)",
    "Archivo PDF Local",
    "Estado Envío",
    "Fecha y Hora Envío"
]

header_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def get_tratamiento(nombre):
    if not nombre:
        return "Estimado(a)"
    nm_lower = nombre.lower()
    if nm_lower.startswith("dra.") or nm_lower.startswith("doctora"):
        return "Estimada"
    if nm_lower.startswith("dr.") or nm_lower.startswith("doctor"):
        return "Estimado"
    femeninos = [
        "erika", "maria", "maría", "nohelia", "zahira", "yesenia", "katherine",
        "elena", "ana", "lady", "veronica", "verónica", "carolina", "julia",
        "liseth", "melissa", "ester", "dahiana", "lina", "celina", "cinthya",
        "fátima", "fatima", "jorgelina", "dorelys", "doris", "fabiola", "mayra",
        "belkis", "katty", "eva", "diana", "andrea", "bertha", "giovana",
        "montserrat", "antonia", "myriam", "laura", "cristina", "araceli",
        "marcelina", "thailing", "alba", "elisa", "angie", "emma", "mariana",
        "isela", "ligia", "mariela", "cassandra", "karina", "wendolyne", "rocio",
        "rocío", "consuelo", "gloria", "patricia", "sandra", "claudia", "monica",
        "mónica", "paola", "martha", "marta", "luz", "carmen", "pilar", "mercedes",
        "guadalupe", "rosario", "concepcion", "concepción", "beatriz", "raquel",
        "ines", "inés", "astrid", "vanessa", "vanesa", "tatiana", "stephanie",
        "stefany", "natalia", "ximena", "sheyla", "nancy", "margarita", "leidy",
        "analia", "analía", "camila", "moncerrath", "paulina"
    ]
    limpio = nm_lower
    for pfx in ["dr.", "dra.", "ing.", "lic.", "prof.", "profesor", "profesora"]:
        if limpio.startswith(pfx):
            limpio = limpio[len(pfx):].strip()
    partes = limpio.split()
    primer = partes[0] if partes else ""
    if primer in femeninos or (len(partes) > 1 and partes[1] in femeninos):
        return f"Estimada {nombre}"
    if primer.endswith("a") and primer not in ["alain", "alaín", "josue", "josué"]:
        return f"Estimada {nombre}"
    return f"Estimado {nombre}"

asunto_fijo = "Certificado Oficial de Asistencia - II Foro de Editores de Revistas Científicas 2026"

for idx, a in enumerate(asistentes, start=2):
    cid = a["id"]
    nom = a["nombre"]
    tratamiento = get_tratamiento(nom)
    email = a.get("email") or ""
    doc = a.get("cedula") or ""
    inst = a.get("institucion") or ""
    pais_ciudad = a.get("pais_ciudad") or a.get("pais") or ""
    qr_url = f"https://juannietoval.github.io/certificaci-n-II-foro/?id={cid}"
    pdf_local = a.get("archivo_carpeta") or f"certificados/asistentes/certificado - {nom}.pdf"
    
    row_vals = [
        cid,
        nom,
        tratamiento,
        email,
        doc if doc else "[NO REGISTRADO]",
        inst,
        pais_ciudad,
        asunto_fijo,
        qr_url,
        pdf_local,
        "PENDIENTE",
        ""
    ]
    ws.append(row_vals)
    
    for col_num in range(1, len(row_vals) + 1):
        cell = ws.cell(row=idx, column=col_num)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if col_num in [1, 3, 4, 11, 12]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if col_num == 11:
                cell.font = Font(name="Calibri", size=10, bold=True, color="0B223D")
        elif col_num == 5:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

ws.row_dimensions[1].height = 28
for r in range(2, len(asistentes) + 2):
    ws.row_dimensions[r].height = 22

col_widths = {
    1: 18,  # ID
    2: 32,  # Nombre
    3: 32,  # Tratamiento
    4: 32,  # Correo
    5: 22,  # Documento
    6: 38,  # Institución
    7: 28,  # País / Ciudad
    8: 38,  # Asunto
    9: 30,  # QR
    10: 38, # Archivo PDF
    11: 16, # Estado
    12: 22  # Fecha/Hora
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

target_1 = os.path.join(BASE, "envio_asistentes.xlsx")
target_2 = r"c:\Users\Lenovo\Desktop\Juan\envio_asistentes.xlsx"

wb.save(target_1)
wb.save(target_2)
print(f"Created {target_1} and {target_2} successfully with {len(asistentes)} asistentes!")
