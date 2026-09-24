import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"
json_path = os.path.join(BASE, "data", "participantes.json")

with open(json_path, "r", encoding="utf-8") as f:
    participantes = json.load(f)

ponentes = [p for p in participantes if p.get("rol") == "PONENTE"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Piloto_Ponentes"
ws.views.sheetView[0].showGridLines = True

headers = [
    "Código Certificado",
    "Nombre Ponente",
    "Tratamiento",
    "Correo Envío Piloto (Prueba)",
    "Correo Real del Ponente",
    "Documento",
    "Institución",
    "Ponencia Magistral Presentada",
    "Eje Temático",
    "Asunto Correo",
    "Enlace Validación Web (QR)",
    "Enlace Grabación (Stream)",
    "Archivo PDF Local",
    "Estado Envío"
]

header_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

# Function to get treatment
def get_tratamiento(nombre):
    nm_lower = nombre.lower()
    if nm_lower.startswith("dr."):
        return f"Estimado {nombre}"
    femeninos = ["nohelia", "zahira", "erika", "yesenia", "katherine", "segunda", "elena", "maría", "maria"]
    primer = nm_lower.split()[0]
    if primer in femeninos or primer.endswith("a"):
        return f"Estimada {nombre}"
    return f"Estimado {nombre}"

LINK_GRABACION = "https://uniminuto0-my.sharepoint.com/personal/felix_duenas_uniminuto_edu/_layouts/15/stream.aspx?id=%2Fpersonal%2Ffelix%5Fduenas%5Funiminuto%5Fedu%2FDocuments%2FGrabaciones%2FII%20FORO%20DE%20EDITORES%20DE%20REVISTAS%20CIENT%C3%8DFICAS%2D20260911%5F080635%2DGrabaci%C3%B3n%20de%20la%20reuni%C3%B3n%2Emp4&ga=1&referrer=StreamWebApp%2EWeb&referrerScenario=AddressBarCopied%2Eview%2E692c7bc5%2Df83b%2D4dbf%2D905d%2D7cb360d59b9e"

missing_count = 0
for idx, p in enumerate(ponentes, start=2):
    cid = p["id"]
    nom = p["nombre"]
    tratamiento = get_tratamiento(nom)
    email_real = p.get("email") or ""
    doc = p.get("cedula") or ""
    inst = p.get("institucion") or ""
    titulo = p.get("titulo") or ""
    eje = p.get("eje") or ""
    qr_url = f"https://juannietoval.github.io/certificaci-n-II-foro/?id={cid}"
    asunto = f"Certificado Oficial de Ponente y Grabación - II Foro de Editores de Revistas Científicas 2026"
    pdf_local = p.get("archivo_carpeta") or f"certificados/ponentes/certificado - {nom}.pdf"
    
    # Destination for pilot: Juan's email
    email_piloto = "juan.nieto2@utp.edu.co"
    
    row_vals = [
        cid,
        nom,
        tratamiento,
        email_piloto,
        email_real if email_real else "[PENDIENTE - NO SUMINISTRADO]",
        doc if doc else "[NO REGISTRADO EN PROGRAMA]",
        inst,
        titulo,
        eje,
        asunto,
        qr_url,
        LINK_GRABACION,
        pdf_local,
        "PENDIENTE"
    ]
    ws.append(row_vals)
    
    # Formatting row
    is_missing = not bool(email_real)
    fill_row = PatternFill(start_color="FFF9E6" if is_missing else "FFFFFF", end_color="FFF9E6" if is_missing else "FFFFFF", fill_type="solid")
    
    for col_num in range(1, len(row_vals) + 1):
        cell = ws.cell(row=idx, column=col_num)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if col_num in [1, 3, 4, 14]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_num in [5, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if is_missing and col_num == 5:
                cell.font = Font(name="Calibri", size=10, bold=True, color="B8860B")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.fill = fill_row

# Set row heights and col widths
ws.row_dimensions[1].height = 28
for r in range(2, len(ponentes) + 2):
    ws.row_dimensions[r].height = 22

col_widths = {
    1: 18,  # ID
    2: 32,  # Nombre
    3: 32,  # Tratamiento
    4: 26,  # Correo Piloto
    5: 28,  # Correo Real
    6: 22,  # Documento
    7: 35,  # Institución
    8: 45,  # Ponencia
    9: 30,  # Eje
    10: 35, # Asunto
    11: 30, # QR
    12: 30, # Stream
    13: 35, # Archivo PDF
    14: 15  # Estado
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

target_file_1 = os.path.join(BASE, "piloto_envio_ponentes.xlsx")
target_file_2 = r"c:\Users\Lenovo\Desktop\Juan\piloto_envio_ponentes.xlsx"

wb.save(target_file_1)
wb.save(target_file_2)

print(f"Created {target_file_1} and {target_file_2} successfully!")
