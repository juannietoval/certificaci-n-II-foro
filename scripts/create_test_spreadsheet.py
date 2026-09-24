# -*- coding: utf-8 -*-
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"
excel_path = os.path.join(BASE, "prueba_envio_dos_certificados.xlsx")

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Prueba_Envio"

# Styles
navy_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
cell_font = Font(name="Calibri", size=10)
center_align = Alignment(horizontal="center", vertical="center")
left_align = Alignment(horizontal="left", vertical="center")
thin_border = Border(
    left=Side(style='thin', color='E0E0E0'),
    right=Side(style='thin', color='E0E0E0'),
    top=Side(style='thin', color='E0E0E0'),
    bottom=Side(style='thin', color='E0E0E0')
)

headers = [
    "Código Certificado",
    "Nombre Completo",
    "Rol",
    "Correo de Envío (Prueba)",
    "Documento Oficial",
    "Institución",
    "Ponencia Magistral / Temática",
    "Eje Temático",
    "Asunto del Correo",
    "Enlace Verificación QR",
    "Enlace Grabación SharePoint",
    "Archivo PDF en Carpeta",
    "Estado Envío"
]

ws.append(headers)
for col_idx in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_idx)
    cell.fill = navy_fill
    cell.font = header_font
    cell.alignment = center_align

link_grabacion = "https://uniminuto0-my.sharepoint.com/personal/felix_duenas_uniminuto_edu/_layouts/15/stream.aspx?id=%2Fpersonal%2Ffelix%5Fduenas%5Funiminuto%5Fedu%2FDocuments%2FGrabaciones%2FII%20FORO%20DE%20EDITORES%20DE%20REVISTAS%20CIENT%C3%8DFICAS%2D20260911%5F080635%2DGrabaci%C3%B3n%20de%20la%20reuni%C3%B3n%2Emp4&ga=1&referrer=StreamWebApp%2EWeb&referrerScenario=AddressBarCopied%2Eview%2E692c7bc5%2Df83b%2D4dbf%2D905d%2D7cb360d59b9e"

rows_data = [
    [
        "FORO26-ASI-012",
        "Juan Esteban Nieto Valencia",
        "ASISTENTE",
        "juan.nieto2@utp.edu.co",
        "C.C. 1.089.602.237",
        "UTP",
        "Gestión editorial en tiempos de inteligencia artificial",
        "General",
        "Certificado Oficial de Asistencia y Grabación - II Foro de Editores de Revistas Científicas 2026",
        "https://juannietoval.github.io/certificaci-n-II-foro/?id=FORO26-ASI-012",
        link_grabacion,
        "certificados/asistentes/certificado - Juan Esteban Nieto Valencia.pdf",
        "PENDIENTE"
    ],
    [
        "FORO26-PON-006",
        "Erika Betancourt",
        "PONENTE",
        "juanestebannietovalencia@gmail.com",
        "",
        "Universidad Tecnológica de Pereira",
        "Presentación de la Revista Miradas",
        "Socialización de Revistas Científicas",
        "Certificado Oficial de Ponente y Grabación - II Foro de Editores de Revistas Científicas 2026",
        "https://juannietoval.github.io/certificaci-n-II-foro/?id=FORO26-PON-006",
        link_grabacion,
        "certificados/ponentes/certificado - Erika Betancourt.pdf",
        "PENDIENTE"
    ]
]

for r_data in rows_data:
    ws.append(r_data)

for r in range(2, ws.max_row + 1):
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=r, column=c)
        cell.font = cell_font
        cell.border = thin_border
        if c in (1, 3, 5, 13):
            cell.alignment = center_align
        else:
            cell.alignment = left_align

ws.row_dimensions[1].height = 26
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = min(max(max_len + 3, 14), 55)

wb.save(excel_path)
print(f"Created test spreadsheet: {excel_path}")
