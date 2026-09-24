# -*- coding: utf-8 -*-
import os
import re
import json
import html
import shutil
import subprocess
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import pymupdf

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"
CHROME = r"C:\Program Files\Google\Chrome\Application\chrome.exe"
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"
BROWSER = CHROME if os.path.exists(CHROME) else EDGE
BASE_WEB_URL = "https://juannietoval.github.io/certificaci-n-II-foro"

# Directories
DIR_OUTPUT_PDF = os.path.join(BASE, "output", "pdf")
DIR_OUTPUT_PREVIEW = os.path.join(BASE, "output", "preview")
DIR_CERTS_BASE = os.path.join(BASE, "certificados")
DIR_CERTS_PONENTES = os.path.join(DIR_CERTS_BASE, "ponentes")
DIR_CERTS_ASISTENTES = os.path.join(DIR_CERTS_BASE, "asistentes")
DIR_TEMPLATES_ASI = os.path.join(BASE, "templates", "asistentes")

for d in [DIR_OUTPUT_PDF, DIR_OUTPUT_PREVIEW, DIR_CERTS_PONENTES, DIR_CERTS_ASISTENTES, DIR_TEMPLATES_ASI]:
    os.makedirs(d, exist_ok=True)

# Clean up test files from output
for folder in [DIR_OUTPUT_PDF, DIR_OUTPUT_PREVIEW]:
    for fname in os.listdir(folder):
        if fname.startswith("TEST_") or fname.startswith("test_"):
            try:
                os.remove(os.path.join(folder, fname))
            except Exception:
                pass

DEFAULT_EVENT_TITLE = "II FORO DE EDITORES DE REVISTAS CIENTÍFICAS"
DEFAULT_EVENT_SUBTITLE = "Gestión editorial en tiempos de inteligencia artificial"
DEFAULT_FECHA = "Viernes 11 de septiembre de 2026"
DEFAULT_HORARIO = "8:00 a. m. a 12:00 m. (hora Colombia)"
DEFAULT_INTENSIDAD = "4 horas"
DEFAULT_MODALIDAD = "Virtual"
DEFAULT_CIUDAD = "Pereira, Colombia"

def clean_val(val, default=""):
    if val is None:
        return default
    text = str(val).strip()
    return html.escape(text)

def safe_filename(name):
    s = re.sub(r'[\/\\:\*\?"<>\|]', '', name)
    s = re.sub(r'\s+', '_', s.strip())
    return s

def title_case_name(name):
    if not name:
        return ""
    lowers = {'de', 'del', 'la', 'las', 'el', 'los', 'y', 'e', 'en', 'da', 'do', 'dos'}
    custom_fixes = {
        'MARIA ESTER GONZALEZ': 'María Ester González',
        'CAMILO CORCHUELO': 'Camilo Corchuelo',
        'ENRIQUE BLANCARTE FUENTES': 'Enrique Blancarte Fuentes',
        'MARIA GUADALUPE RODRIGUEZ OLIVA': 'María Guadalupe Rodríguez Oliva',
        'CORINA ECHAVARRIA': 'Corina Echavarría',
        'BERTHA JACQUELINE CONTLA RAMIREZ': 'Bertha Jacqueline Contla Ramírez',
        'MARCELINA SOYDETH JIMENEZ AVILA': 'Marcelina Soydeth Jiménez Ávila',
        'THAILING NUNEZ BETANCOURT': 'Thailing Núñez Betancourt',
        'BAYRON STEVEN LOPEZ JAUREUI': 'Bayron Steven López Jáuregui',
        'Fabiola Nez Pastrana': 'Fabiola Núñez Pastrana',
        'Fabiola N?ez Pastrana': 'Fabiola Núñez Pastrana'
    }
    clean_n = re.sub(r'\s+', ' ', name).strip()
    if clean_n in custom_fixes:
        return custom_fixes[clean_n]
    if clean_n.upper() in custom_fixes:
        return custom_fixes[clean_n.upper()]
    if clean_n.isupper():
        words = clean_n.split()
        res = []
        for i, w in enumerate(words):
            wl = w.lower()
            if i > 0 and wl in lowers:
                res.append(wl)
            else:
                res.append(w.capitalize())
        return " ".join(res)
    return clean_n

def format_doc_asistente(tipo_doc, doc_str):
    if not doc_str:
        return ""
    s = str(doc_str).strip()
    if s.upper() in ('NA', 'N/A', 'NONE', 'NO', '-', '.', ''):
        return ""
    upper_s = s.upper()
    if upper_s.startswith("IDMEX"):
        return f"INE {s}"
    prefixes = ("C.C.", "CC", "D.I.", "DI", "DNI", "PASAPORTE", "PAS.", "PAS", "C.E.", "CE", "C.I.", "CI", "INE", "CURP", "RUT")
    for pfx in prefixes:
        if upper_s.startswith(pfx):
            return s
    upper_td = str(tipo_doc or '').upper()
    if 'CIUDADAN' in upper_td or 'C.C' in upper_td:
        if s.isdigit():
            num = int(s)
            return f"C.C. {num:,}".replace(",", ".")
        return f"C.C. {s}"
    elif 'EXTRANJER' in upper_td or 'C.E' in upper_td:
        return f"C.E. {s}"
    elif 'IDENTIDAD' in upper_td or 'C.I' in upper_td or 'DNI' in upper_td:
        if 'DNI' in upper_td:
            return f"DNI {s}"
        return f"C.I. {s}"
    elif 'PASAPORTE' in upper_td:
        return f"Pasaporte {s}"
    elif 'INE' in upper_td:
        return f"INE {s}"
    if s.isdigit():
        num = int(s)
        return f"C.C. {num:,}".replace(",", ".")
    return f"D.I. {s}"

def main():
    print("=== STARTING FULL PRODUCTION PIPELINE ===")

    # 1. Load Templates
    with open(os.path.join(BASE, "templates", "master_template.html"), "r", encoding="utf-8") as f:
        template_ponente = f.read()
    with open(os.path.join(BASE, "templates", "master_asistente_template.html"), "r", encoding="utf-8") as f:
        template_asistente = f.read()

    # 2. Load Ponentes
    data_ponentes_path = os.path.join(BASE, "data", "participantes.json")
    with open(data_ponentes_path, "r", encoding="utf-8") as f:
        all_participantes = json.load(f)
    ponentes = [p for p in all_participantes if p.get("rol", "").upper() == "PONENTE"]
    print(f"Loaded {len(ponentes)} confirmed Ponentes.")

    # 3. Load Asistentes Raw
    with open(os.path.join(BASE, "data", "asistentes_raw.json"), "r", encoding="utf-8") as f:
        raw_asistentes = json.load(f)
    print(f"Loaded {len(raw_asistentes)} raw attendees to process.")

    # 4. Copy Ponentes to certificados/ponentes/
    print("\n--- ORGANIZING PONENTES CERTIFICATES ---")
    for p in ponentes:
        cert_id = p["id"]
        nombre = p["nombre"]
        clean_name = safe_filename(nombre)
        src_pdf = os.path.join(DIR_OUTPUT_PDF, f"{cert_id}_{clean_name}.pdf")
        dest_pdf = os.path.join(DIR_CERTS_PONENTES, f"certificado - {nombre}.pdf")
        if os.path.exists(src_pdf):
            shutil.copy2(src_pdf, dest_pdf)
            print(f"  [OK] Copied {src_pdf} -> {dest_pdf}")
        else:
            print(f"  [WARN] Source PDF not found: {src_pdf}")

    # 5. Process and Generate Asistentes
    print(f"\n--- GENERATING {len(raw_asistentes)} ASISTENTES CERTIFICATES ---")
    asistentes_records = []
    
    for idx, a in enumerate(raw_asistentes):
        cert_id = f"FORO26-ASI-{idx+1:03d}"
        nombre = title_case_name(a["nombre"])
        clean_name = safe_filename(nombre)
        doc_line = format_doc_asistente(a["tipo_doc"], a["doc"])
        institucion = a.get("institucion", "").strip()
        pais_ciudad = a.get("pais_ciudad", "").strip()
        email = a.get("email", "").strip()
        
        ciudad_disp = f"{DEFAULT_CIUDAD} (Modalidad {DEFAULT_MODALIDAD})"
        location_date = f"{ciudad_disp} &mdash; {DEFAULT_FECHA}"
        qr_url = f"{BASE_WEB_URL}/?id={cert_id}"
        
        # Build HTML
        cert_html = template_asistente
        cert_html = cert_html.replace("{{ID}}", cert_id)
        cert_html = cert_html.replace("{{NOMBRE}}", clean_val(nombre))
        cert_html = cert_html.replace("{{DOCUMENTO_LINE}}", clean_val(doc_line))
        cert_html = cert_html.replace("{{ROL}}", "ASISTENTE")
        cert_html = cert_html.replace("{{INTENSIDAD}}", DEFAULT_INTENSIDAD)
        cert_html = cert_html.replace("{{HORARIO}}", DEFAULT_HORARIO)
        
        # Affiliation line
        if institucion and pais_ciudad:
            aff_line = f'{clean_val(institucion)} &bull; {clean_val(pais_ciudad)}'
        elif institucion:
            aff_line = clean_val(institucion)
        elif pais_ciudad:
            aff_line = clean_val(pais_ciudad)
        else:
            aff_line = ''
            
        if aff_line:
            aff_block = f'<div class="affiliation">{aff_line}</div>'
        else:
            aff_block = ''
        cert_html = cert_html.replace("{{AFFILIATION_BLOCK}}", aff_block)
        cert_html = cert_html.replace("{{LOCATION_DATE}}", location_date)
        cert_html = cert_html.replace("{{COORDINACION}}", "Comité Organizador")
        cert_html = cert_html.replace("{{QR_URL}}", qr_url)
        
        # Write HTML
        html_file = os.path.join(DIR_TEMPLATES_ASI, f"{cert_id}.html")
        with open(html_file, "w", encoding="utf-8") as f_out:
            f_out.write(cert_html)
            
        # Target PDF paths
        pdf_official_rel = f"output/pdf/{cert_id}_{clean_name}.pdf"
        pdf_official_abs = os.path.join(BASE, pdf_official_rel)
        preview_rel = f"output/preview/{cert_id}_preview.png"
        preview_abs = os.path.join(BASE, preview_rel)
        pdf_user_facing = os.path.join(DIR_CERTS_ASISTENTES, f"certificado - {nombre}.pdf")
        
        # Render PDF via headless browser
        cmd_pdf = [
            BROWSER,
            "--headless",
            "--disable-gpu",
            "--no-pdf-header-footer",
            f"--print-to-pdf={pdf_official_abs}",
            f"file:///{html_file}"
        ]
        subprocess.run(cmd_pdf, capture_output=True, text=True)
        
        # Copy to certificados/asistentes/certificado - [Nombre].pdf
        if os.path.exists(pdf_official_abs):
            shutil.copy2(pdf_official_abs, pdf_user_facing)
            
        # Render PNG preview (150 DPI)
        try:
            doc = pymupdf.open(pdf_official_abs)
            page = doc[0]
            pix = page.get_pixmap(dpi=150)
            pix.save(preview_abs)
            doc.close()
        except Exception as e:
            print(f"  [ERROR] Preview failed for {cert_id}: {e}")
            
        # Build participant dictionary
        record = {
            "id": cert_id,
            "nombre": nombre,
            "cedula": doc_line,
            "tipo_doc": a.get("tipo_doc", ""),
            "doc_raw": a.get("doc", ""),
            "rol": "ASISTENTE",
            "tema_central": DEFAULT_EVENT_SUBTITLE,
            "intensidad": DEFAULT_INTENSIDAD,
            "modalidad": DEFAULT_MODALIDAD,
            "institucion": institucion,
            "pais_ciudad": pais_ciudad,
            "fecha": DEFAULT_FECHA,
            "horario": DEFAULT_HORARIO,
            "ciudad": DEFAULT_CIUDAD,
            "coordinacion": "Comité Organizador",
            "email": email,
            "pdf": pdf_official_rel,
            "preview": preview_rel,
            "archivo_carpeta": f"certificados/asistentes/certificado - {nombre}.pdf"
        }
        asistentes_records.append(record)
        
        if (idx + 1) % 10 == 0 or (idx + 1) == len(raw_asistentes):
            print(f"  Processed {idx + 1}/{len(raw_asistentes)}: {cert_id} - {nombre}")

    # 6. Build Consolidated Participantes List (12 ponentes + 102 asistentes = 114)
    print("\n--- UPDATING CONSOLIDATED JSON DATABASE ---")
    consolidated_list = []
    for p in ponentes:
        p_copy = dict(p)
        p_copy["archivo_carpeta"] = f"certificados/ponentes/certificado - {p['nombre']}.pdf"
        p_copy["pdf"] = f"output/pdf/{p['id']}_{safe_filename(p['nombre'])}.pdf"
        p_copy["preview"] = f"output/preview/{p['id']}_preview.png"
        consolidated_list.append(p_copy)
    consolidated_list.extend(asistentes_records)

    with open(os.path.join(BASE, "data", "participantes.json"), "w", encoding="utf-8") as f:
        json.dump(consolidated_list, f, ensure_ascii=False, indent=2)
    print(f"Successfully saved {len(consolidated_list)} participants in data/participantes.json")

    # 7. Update index.html
    print("\n--- UPDATING WEB VALIDATION PORTAL (index.html) ---")
    index_path = os.path.join(BASE, "index.html")
    with open(index_path, "r", encoding="utf-8") as f:
        index_html = f.read()
    json_str = json.dumps(consolidated_list, ensure_ascii=False, indent=6)
    pattern = r"const DEFAULT_CERTS = \[.*?\];"
    replacement = f"const DEFAULT_CERTS = {json_str};"
    new_index_html = re.sub(pattern, replacement, index_html, flags=re.DOTALL)
    with open(index_path, "w", encoding="utf-8") as f:
        f.write(new_index_html)
    print(f"index.html updated with {len(consolidated_list)} verified records.")

    # 8. Create Consolidated Excel Spreadsheet
    print("\n--- CREATING CONSOLIDATED EXCEL DATABASE ---")
    create_excel_database(ponentes, asistentes_records)

    print("\n=== PIPELINE COMPLETED SUCCESSFULLY ===")

def create_excel_database(ponentes, asistentes):
    excel_path = os.path.join(BASE, "registro_certificados_foro_2026.xlsx")
    wb = openpyxl.Workbook()
    # remove default sheet
    wb.remove(wb.active)

    # Styles
    navy_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
    gold_fill = PatternFill(start_color="8C6D3B", end_color="8C6D3B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )

    # Sheet 1: Ponentes
    ws_pon = wb.create_sheet(title="Ponentes")
    headers_pon = [
        "Código Certificado", "Nombre Completo", "Cédula / Documento", "Institución",
        "País", "Ponencia Magistral Presentada", "Eje Temático", "Intensidad",
        "Enlace Verificación QR", "Archivo PDF en Carpeta"
    ]
    ws_pon.append(headers_pon)
    for col_idx, h in enumerate(headers_pon, 1):
        cell = ws_pon.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = center_align

    for p in ponentes:
        row_data = [
            p.get("id", ""),
            p.get("nombre", ""),
            p.get("cedula", ""),
            p.get("institucion", ""),
            p.get("pais", ""),
            p.get("titulo", ""),
            p.get("eje", ""),
            p.get("intensidad", "4 horas"),
            f"{BASE_WEB_URL}/?id={p.get('id', '')}",
            f"certificados/ponentes/certificado - {p.get('nombre', '')}.pdf"
        ]
        ws_pon.append(row_data)

    # Format data rows
    for r in range(2, ws_pon.max_row + 1):
        for c in range(1, len(headers_pon) + 1):
            cell = ws_pon.cell(row=r, column=c)
            cell.font = normal_font
            cell.border = thin_border
            if c in (1, 3, 5, 8):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Sheet 2: Asistentes
    ws_asi = wb.create_sheet(title="Asistentes")
    headers_asi = [
        "Código Certificado", "Nombre Completo", "Documento Oficial", "Institución",
        "País / Ciudad", "Temática Central del Evento", "Intensidad", "Correo Electrónico",
        "Enlace Verificación QR", "Archivo PDF en Carpeta"
    ]
    ws_asi.append(headers_asi)
    for col_idx, h in enumerate(headers_asi, 1):
        cell = ws_asi.cell(row=1, column=col_idx)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = center_align

    for a in asistentes:
        row_data = [
            a.get("id", ""),
            a.get("nombre", ""),
            a.get("cedula", ""),
            a.get("institucion", ""),
            a.get("pais_ciudad", ""),
            a.get("tema_central", DEFAULT_EVENT_SUBTITLE),
            a.get("intensidad", "4 horas"),
            a.get("email", ""),
            f"{BASE_WEB_URL}/?id={a.get('id', '')}",
            a.get("archivo_carpeta", "")
        ]
        ws_asi.append(row_data)

    for r in range(2, ws_asi.max_row + 1):
        for c in range(1, len(headers_asi) + 1):
            cell = ws_asi.cell(row=r, column=c)
            cell.font = normal_font
            cell.border = thin_border
            if c in (1, 3, 7):
                cell.alignment = center_align
            else:
                cell.alignment = left_align

    # Auto-adjust column widths
    for ws in [ws_pon, ws_asi]:
        ws.row_dimensions[1].height = 26
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)

    wb.save(excel_path)
    print(f"Consolidated Excel saved to: {excel_path}")

if __name__ == "__main__":
    main()
