import os
import openpyxl

files = [
    r"c:\Users\Lenovo\Desktop\Juan\piloto_envio_ponentes.xlsx",
    r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados\piloto_envio_ponentes.xlsx",
    r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados\prueba_envio_dos_certificados.xlsx"
]

for f in files:
    if os.path.exists(f):
        wb = openpyxl.load_workbook(f)
        for sname in wb.sheetnames:
            ws = wb[sname]
            for row in ws.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        if " y Grabación" in cell.value:
                            cell.value = cell.value.replace(" y Grabación", "")
                            print(f"Updated cell in {f} [{sname}]: {cell.value}")
                        elif " y Grabacion" in cell.value:
                            cell.value = cell.value.replace(" y Grabacion", "")
                            print(f"Updated cell in {f} [{sname}]: {cell.value}")
        wb.save(f)
        print(f"Saved {f} successfully!")
