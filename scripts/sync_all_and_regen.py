import os
import json
import re
import openpyxl
from generator import generate_all

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"

def sync_index_html():
    json_path = os.path.join(BASE, "data", "participantes.json")
    with open(json_path, "r", encoding="utf-8") as f:
        participantes = json.load(f)

    html_path = os.path.join(BASE, "index.html")
    with open(html_path, "r", encoding="utf-8") as f:
        html_content = f.read()

    # Locate DEFAULT_CERTS
    start_marker = "const DEFAULT_CERTS = ["
    end_marker = "];\n\n    let dynamicList = null;"
    if end_marker not in html_content:
        end_marker = "];"
    
    start_idx = html_content.find(start_marker)
    if start_idx == -1:
        raise ValueError("Could not find const DEFAULT_CERTS in index.html")
    
    end_idx = html_content.find("];", start_idx)
    if end_idx == -1:
        raise ValueError("Could not find closing ]; in index.html")
    
    json_formatted = json.dumps(participantes, ensure_ascii=False, indent=6)
    new_certs_block = f"const DEFAULT_CERTS = {json_formatted}"
    
    updated_html = html_content[:start_idx] + new_certs_block + html_content[end_idx + 1:]
    
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(updated_html)
    print("Successfully synchronized index.html with participantes.json!")

def sync_excel_master():
    json_path = os.path.join(BASE, "data", "participantes.json")
    with open(json_path, "r", encoding="utf-8") as f:
        participantes = json.load(f)

    p_by_id = {p["id"]: p for p in participantes}

    excel_path = os.path.join(BASE, "registro_certificados_foro_2026.xlsx")
    wb = openpyxl.load_workbook(excel_path)

    # 1. Update Ponentes
    ws_pon = wb["Ponentes"]
    for row in ws_pon.iter_rows(min_row=2):
        cert_id = row[0].value
        if cert_id in p_by_id:
            p = p_by_id[cert_id]
            row[2].value = p.get("cedula", "")
            row[8].value = p.get("email", "")

    # 2. Update Asistentes
    ws_asi = wb["Asistentes"]
    for row in ws_asi.iter_rows(min_row=2):
        cert_id = row[0].value
        if cert_id in p_by_id:
            p = p_by_id[cert_id]
            row[2].value = p.get("cedula", "")
            row[7].value = p.get("email", "")

    wb.save(excel_path)
    print("Successfully synchronized registro_certificados_foro_2026.xlsx!")

if __name__ == "__main__":
    sync_index_html()
    sync_excel_master()
    print("\nRegenerating all certificates with updated metadata...")
    generate_all()
    print("\nAll done!")
