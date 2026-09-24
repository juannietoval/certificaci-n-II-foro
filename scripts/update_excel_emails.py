# -*- coding: utf-8 -*-
import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

excel_path = r'c:\Users\Lenovo\Desktop\Juan\sistema_certificados\registro_certificados_foro_2026.xlsx'
json_path = r'c:\Users\Lenovo\Desktop\Juan\sistema_certificados\data\participantes.json'

with open(json_path, 'r', encoding='utf-8') as f:
    participantes = json.load(f)

ponentes = [p for p in participantes if p.get('rol') == 'PONENTE']

print("--- PONENTES EMAILS ---")
for p in ponentes:
    print(f"{p['id']} | {p['nombre']} | Email: {p.get('email', '')}")

# Update Excel ws_pon to include Correo Electrónico
wb = openpyxl.load_workbook(excel_path)
ws_pon = wb["Ponentes"]

# Check if Correo Electrónico is in headers
headers = [ws_pon.cell(1, col).value for col in range(1, ws_pon.max_column + 1)]
if "Correo Electrónico" not in headers:
    # Insert Correo Electrónico before Enlace Verificación QR
    col_idx = headers.index("Enlace Verificación QR") + 1
    ws_pon.insert_cols(col_idx)
    ws_pon.cell(1, col_idx).value = "Correo Electrónico"
    
    # Styles
    navy_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    cell_font = Font(name="Calibri", size=10)
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    
    ws_pon.cell(1, col_idx).fill = navy_fill
    ws_pon.cell(1, col_idx).font = header_font
    ws_pon.cell(1, col_idx).alignment = center_align
    
    # Fill in emails
    p_map = {p['id']: p.get('email', '') for p in ponentes}
    for r in range(2, ws_pon.max_row + 1):
        cid = ws_pon.cell(r, 1).value
        cell = ws_pon.cell(r, col_idx)
        cell.value = p_map.get(cid, '')
        cell.font = cell_font
        cell.border = thin_border
        cell.alignment = left_align
        
    for col in ws_pon.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws_pon.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 55)
        
    wb.save(excel_path)
    print("Added 'Correo Electrónico' column to Ponentes sheet in Excel!")
else:
    print("'Correo Electrónico' already in Ponentes sheet.")
