import os
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE = r"c:\Users\Lenovo\Desktop\Juan\sistema_certificados"
json_path = os.path.join(BASE, "data", "participantes.json")

with open(json_path, "r", encoding="utf-8") as f:
    participantes = json.load(f)

ponentes = [p for p in participantes if p.get("rol") == "PONENTE"]

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Piloto_Ponentes"
ws.views.sheetView[0].showGridLines = True

headers = [
    "Código Certificado",
    "Nombre Ponente",
    "Tratamiento",
    "Correo Envío Piloto (Prueba)",
    "Correo Real del Ponente",
    "Documento",
    "Institución",
    "Ponencia Magistral Presentada",
    "Eje Temático",
    "Asunto Correo",
    "Enlace Validación Web (QR)",
    "Archivo PDF Local",
    "Estado Envío"
]

header_fill = PatternFill(start_color="0B223D", end_color="0B223D", fill_type="solid")
header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
border_thin = Border(
    left=Side(style='thin', color='D0D7DE'),
    right=Side(style='thin', color='D0D7DE'),
    top=Side(style='thin', color='D0D7DE'),
    bottom=Side(style='thin', color='D0D7DE')
)

ws.append(headers)
for col_num in range(1, len(headers) + 1):
    cell = ws.cell(row=1, column=col_num)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

def get_tratamiento(nombre):
    nm_lower = nombre.lower()
    if nm_lower.startswith("dr."):
        return f"Estimado {nombre}"
    femeninos = ["nohelia", "zahira", "erika", "yesenia", "katherine", "segunda", "elena", "maría", "maria"]
    primer = nm_lower.split()[0]
    if primer in femeninos or primer.endswith("a"):
        return f"Estimada {nombre}"
    return f"Estimado {nombre}"

for idx, p in enumerate(ponentes, start=2):
    cid = p["id"]
    nom = p["nombre"]
    tratamiento = get_tratamiento(nom)
    email_real = p.get("email") or ""
    doc = p.get("cedula") or ""
    inst = p.get("institucion") or ""
    titulo = p.get("titulo") or ""
    eje = p.get("eje") or ""
    qr_url = f"https://juannietoval.github.io/certificaci-n-II-foro/?id={cid}"
    asunto = "Certificado Oficial de Ponente - II Foro de Editores de Revistas Científicas 2026"
    pdf_local = p.get("archivo_carpeta") or f"certificados/ponentes/certificado - {nom}.pdf"
    
    # Destination for pilot: Juan's email
    email_piloto = "juan.nieto2@utp.edu.co"
    
    row_vals = [
        cid,
        nom,
        tratamiento,
        email_piloto,
        email_real if email_real else "[PENDIENTE - NO SUMINISTRADO]",
        doc if doc else "[NO REGISTRADO EN PROGRAMA]",
        inst,
        titulo,
        eje,
        asunto,
        qr_url,
        pdf_local,
        "PENDIENTE"
    ]
    ws.append(row_vals)
    
    is_missing = not bool(email_real)
    fill_row = PatternFill(start_color="FFF9E6" if is_missing else "FFFFFF", end_color="FFF9E6" if is_missing else "FFFFFF", fill_type="solid")
    
    for col_num in range(1, len(row_vals) + 1):
        cell = ws.cell(row=idx, column=col_num)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if col_num in [1, 3, 4, 13]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        elif col_num in [5, 6]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
            if is_missing and col_num == 5:
                cell.font = Font(name="Calibri", size=10, bold=True, color="B8860B")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")
        cell.fill = fill_row

ws.row_dimensions[1].height = 28
for r in range(2, len(ponentes) + 2):
    ws.row_dimensions[r].height = 22

col_widths = {
    1: 18,  # ID
    2: 32,  # Nombre
    3: 32,  # Tratamiento
    4: 26,  # Correo Piloto
    5: 28,  # Correo Real
    6: 22,  # Documento
    7: 35,  # Institución
    8: 45,  # Ponencia
    9: 30,  # Eje
    10: 38, # Asunto
    11: 30, # QR
    12: 35, # Archivo PDF
    13: 15  # Estado
}

for col_idx, width in col_widths.items():
    ws.column_dimensions[get_column_letter(col_idx)].width = width

target_file_1 = os.path.join(BASE, "piloto_envio_ponentes.xlsx")
target_file_2 = r"c:\Users\Lenovo\Desktop\Juan\piloto_envio_ponentes.xlsx"

wb.save(target_file_1)
wb.save(target_file_2)
print("Updated piloto_envio_ponentes.xlsx with 13 ponentes (without recording link)!")

# Update Master Excel registro_certificados_foro_2026.xlsx
master_file = os.path.join(BASE, "registro_certificados_foro_2026.xlsx")
wb_m = openpyxl.Workbook()

# Sheet 1: Ponentes
ws_m_pon = wb_m.active
ws_m_pon.title = "Ponentes"
ws_m_pon.views.sheetView[0].showGridLines = True
headers_pon = [
    "Código Certificado", "Nombre Completo", "Cédula / Documento", "Institución",
    "País", "Ponencia Magistral Presentada", "Eje Temático", "Intensidad",
    "Correo Electrónica", "Enlace Verificación QR", "Archivo PDF en Carpeta"
]
ws_m_pon.append(headers_pon)
for c in range(1, len(headers_pon) + 1):
    cell = ws_m_pon.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

for idx, p in enumerate(ponentes, start=2):
    r_vals = [
        p["id"],
        p["nombre"],
        p.get("cedula") or "",
        p.get("institucion") or "",
        p.get("pais") or "Colombia",
        p.get("titulo") or "",
        p.get("eje") or "",
        p.get("intensidad") or "4 horas",
        p.get("email") or "",
        f"https://juannietoval.github.io/certificaci-n-II-foro/?id={p['id']}",
        p.get("archivo_carpeta") or ""
    ]
    ws_m_pon.append(r_vals)
    for c in range(1, len(r_vals) + 1):
        cell = ws_m_pon.cell(row=idx, column=c)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if c in [1, 3, 5, 8, 9]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

# Sheet 2: Asistentes
ws_m_asi = wb_m.create_sheet(title="Asistentes")
ws_m_asi.views.sheetView[0].showGridLines = True
headers_asi = [
    "Código Certificado", "Nombre Completo", "Documento Oficial", "Institución",
    "País / Ciudad", "Temática Central del Evento", "Intensidad",
    "Correo Electrónico", "Enlace Verificación QR", "Archivo PDF en Carpeta"
]
ws_m_asi.append(headers_asi)
for c in range(1, len(headers_asi) + 1):
    cell = ws_m_asi.cell(row=1, column=c)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

asistentes = [p for p in participantes if p.get("rol") == "ASISTENTE"]
for idx, a in enumerate(asistentes, start=2):
    r_vals = [
        a["id"],
        a["nombre"],
        a.get("cedula") or "",
        a.get("institucion") or "",
        a.get("pais_ciudad") or a.get("pais") or "",
        a.get("tema_central") or "Gestión editorial en tiempos de inteligencia artificial",
        a.get("intensidad") or "4 horas",
        a.get("email") or "",
        f"https://juannietoval.github.io/certificaci-n-II-foro/?id={a['id']}",
        a.get("archivo_carpeta") or ""
    ]
    ws_m_asi.append(r_vals)
    for c in range(1, len(r_vals) + 1):
        cell = ws_m_asi.cell(row=idx, column=c)
        cell.font = Font(name="Calibri", size=10)
        cell.border = border_thin
        if c in [1, 3, 7, 8]:
            cell.alignment = Alignment(horizontal="center", vertical="center")
        else:
            cell.alignment = Alignment(horizontal="left", vertical="center")

wb_m.save(master_file)
print(f"Updated master Excel {master_file} successfully!")
